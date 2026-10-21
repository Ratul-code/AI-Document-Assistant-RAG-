from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from app.models.document_page import TextSource
from app.services.document.parser import parse_document


def test_docx_splits_on_heading_1(tiny_docx_path: Path) -> None:
    pages = parse_document(
        tiny_docx_path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "tiny.docx",
    )
    assert len(pages) == 3
    titles = [p.raw_text.split("\n\n")[0] for p in pages]
    assert titles == ["Introduction", "Methods", "Results"]
    for page in pages:
        assert page.text_source == TextSource.raw
        assert page.parser_name == "python-docx"
        assert page.page_metadata["split_strategy"] == "heading1"


def test_docx_no_heading_1_returns_single_page(flat_docx_path: Path) -> None:
    pages = parse_document(
        flat_docx_path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "flat.docx",
    )
    assert len(pages) == 1
    assert pages[0].text_source == TextSource.raw
    assert pages[0].page_metadata["split_strategy"] == "single"
    assert "Plain first paragraph." in pages[0].raw_text


def test_docx_renders_table_as_markdown(table_docx_path: Path) -> None:
    pages = parse_document(
        table_docx_path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "table.docx",
    )
    assert len(pages) == 1
    raw = pages[0].raw_text
    # Surrounding paragraphs are present.
    assert "Intro paragraph above the table." in raw
    assert "Outro paragraph below the table." in raw
    # Markdown table is rendered (header, separator, body rows).
    assert "| Name | Value |" in raw
    assert "| --- |" in raw
    assert "| alpha | 1 |" in raw
    assert "| beta | 2 |" in raw
    # Metadata captures the table structure.
    tables = pages[0].page_metadata["tables"]
    assert len(tables) == 1
    assert tables[0]["row_count"] == 3
    assert tables[0]["col_count"] == 2
    assert tables[0]["markdown"] in raw


def test_docx_pipe_in_cell_is_escaped(tmp_path: Path) -> None:
    """A literal '|' inside a cell must be escaped so the table stays parseable."""
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A | B"
    table.cell(0, 1).text = "C"
    table.cell(1, 0).text = "D"
    table.cell(1, 1).text = "E"
    path = tmp_path / "pipes.docx"
    doc.save(str(path))
    pages = parse_document(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pipes.docx",
    )
    raw = pages[0].raw_text
    assert r"A \| B" in raw
    # Cell text is still readable after the escape.
    assert "C" in raw
